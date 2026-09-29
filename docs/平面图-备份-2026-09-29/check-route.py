# -*- coding: utf-8 -*-
"""自检渲染：把 Node 算出的最优路线画到底图上，供我肉眼核对（压不压货架、有没有斜穿）。
**独立性**：可走性直接回读 Excel 的填充色（theme3）判定，不复用 route-core 的结论 ——
若 route-core 建模有误，这里会报「非法格」。
坐标与 render-excel.py 一致：x = 190+(col-1)*13，y = 66+(row-1)*13（1 单元格 = 13px）
"""
import io, json, sys
import openpyxl
from PIL import Image, ImageDraw, ImageFont

BASE = r'E:\claude\parcel-SPU\.devtools\floorplan-excel.png'
SAMPLE = sys.argv[1] if len(sys.argv) > 1 else r'E:\claude\parcel-SPU\.devtools\route-sample.json'
SRC = r'E:\claude\parcel-SPU\docs\海大快递站平面图-2026-09-29-用户精确版.xlsx'
OUT = r'E:\claude\parcel-SPU\.devtools\floorplan-check.png' if len(sys.argv) <= 1 \
    else r'E:\claude\parcel-SPU\.devtools\floorplan-check-page.png'
L_LANE, T, CELL = 190, 66, 13

sample = json.load(io.open(SAMPLE, encoding='utf-8'))
wb = openpyxl.load_workbook(SRC)
ws = wb['Sheet1']


def fillkey(c):
    f = c.fill
    if f is None or f.fill_type is None or f.fgColor is None:
        return None
    fg = f.fgColor
    if fg.type == 'rgb' and fg.rgb:
        return str(fg.rgb)
    if fg.type == 'theme':
        return 'theme%d' % fg.theme
    return None


gate_cells = set()
for mr in ws.merged_cells.ranges:
    v = ws.cell(row=mr.min_row, column=mr.min_col).value
    if v and ('闸机' in str(v) or '出口' in str(v)):
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                gate_cells.add((r, c))


def walkable(r, c):
    return (r, c) in gate_cells or fillkey(ws.cell(row=r, column=c)) == 'theme3'


cx = lambda c: L_LANE + (c - 1) * CELL + CELL / 2
cy = lambda r: T + (r - 1) * CELL + CELL / 2

img = Image.open(BASE).convert('RGB')
dr = ImageDraw.Draw(img)
f = ImageFont.truetype(r'C:\Windows\Fonts\msyhbd.ttc', 15)
fs = ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc', 13)

bad = []
for leg in sample['legs']:
    cells = leg.get('cells') or []
    for i, (r, c) in enumerate(cells):
        if not walkable(r, c):
            bad.append(('非法格', leg['from'] + '→' + leg['to'], i, (r, c)))
        if i:
            pr, pc = cells[i - 1]
            if abs(pr - r) + abs(pc - c) != 1:
                bad.append(('不相邻', leg['from'] + '→' + leg['to'], (pr, pc), (r, c)))
    if len(cells) >= 2:
        dr.line([(cx(c), cy(r)) for (r, c) in cells], fill=(210, 25, 25), width=3, joint='curve')

visitNo = {code: i + 1 for i, code in enumerate(sample['order'])}
for p in sample['picks']:
    x, y = cx(p['col']), cy(p['row'])
    dr.ellipse([x - 11, y - 11, x + 11, y + 11], fill=(31, 111, 235), outline='white', width=2)
    dr.text((x, y), str(visitNo.get(p['code'], '?')), font=f, fill='white', anchor='mm')

ex, ey = cx(sample['entrance']['col']), cy(sample['entrance']['row'])
dr.ellipse([ex - 12, ey - 12, ex + 12, ey + 12], fill=(21, 128, 61), outline='white', width=2)
dr.text((ex, ey), '入', font=f, fill='white', anchor='mm')

for (r0, r1, c0, c1) in sample['gateSpan']:
    dr.rectangle([L_LANE + (c0 - 1) * CELL, T + (r0 - 1) * CELL, L_LANE + c1 * CELL, T + r1 * CELL],
                 outline=(185, 28, 28), width=2)

dr.text((14, 62), '自检：红线＝模型算出的最优路线　蓝点＝取件顺序　绿点＝入口　红框＝闸机带　浅蓝＝Excel 填充的通道',
        font=fs, fill='#7f1d1d')
img.save(OUT)
print('检查图:', OUT, img.size)
print('顺序:', ' → '.join(sample['order']), '→ 闸机')
print('总距离: %.1f 瓷砖' % sample['total'])
print('非法格/斜穿:', len(bad))
for b in bad[:10]:
    print('   ', b)
